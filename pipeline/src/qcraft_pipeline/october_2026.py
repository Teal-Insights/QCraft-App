"""Offline October Current rebuild. Existing vintages and engine code are inputs.

Run: python -m qcraft_pipeline.october_2026 --source-dir <frozen-cache> --out <new-dir>
No network access, partial rosters, implicit vintage changes or output overwrites.
"""

import argparse
import csv
import hashlib
import json
import math
import re
import shutil
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl
from qcraft_engine.horizon import resolve_horizon

from qcraft_pipeline import config, emit, fetch, weo

RETAINED_FROM = "weo-2026-04-full-horizon-v1"
SOURCE_VINTAGE = "weo-2026-10"
REVISION = "weo-2026-10-full-horizon-v1"
POLICY = "current-full-weo-v1"
LABEL = "WEO October 2026 + UN WPP 2024"
INPUT_KEYS = (
    "iso3c",
    "country",
    "macrofiscal",
    "demography",
    "productivity",
    "climate",
)
INDICATORS = tuple(config.WEO_INDICATORS)
OLD_TREES = {
    "weo-2024-10": "171362191850e3a296062eb5e229e2b69e810a264e02a60a5f488d5e8c54e2e3",
    "weo-2026-04": "c9bf51819711000d59750cca64ce2ea52a69086793d367ca0615753f56aa204a",
    RETAINED_FROM: "b4c9175160add1d0ddbed5819839625538b846c9973ccc63021fe2db398b6103",
}
SOURCE_FILES = {
    "source_receipts.json": (
        "ad573d844dc4ae80335197d75f609d0fcc87306f511aa55903b2b738abb1270c"
    ),
    "weo_oct2026_raw.csv": (
        "46d614d31b38aaf71fd825d9d24e8e6a449ef998558b94f21e65c7dcdedffd24"
    ),
    "WEOOct2026all.xlsx": (
        "2aa4f70a05931197b41ed7e0202dc56eee8d816858ac14521e4f8f5dbc38696c"
    ),
    "cl_country.json": (
        "c2565a8dc940223906a4c5c67112443715e2e16bca17a885024528f747595c8a"
    ),
}
RAW_CSV_PATH = config.cache_dir() / "october-2026" / "weo_oct2026_raw.csv"
NUMBER = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][+-]?\d+)?\Z")
WORKBOOK_PREFIX = """DATASET SERIES_CODE COUNTRY.ID COUNTRY INDICATOR.ID INDICATOR
INDICATOR.Description FREQUENCY SCALE UNIT COUNTRY_UPDATE_DATE PUBLICATION_DATE
UPDATE_DATE METHODOLOGY.ID METHODOLOGY_NOTES LATEST_ACTUAL_ANNUAL_DATA
HISTORICAL_DATA_SOURCE BASE_YEAR START_END_MONTHS_OF_REPORTING_YEAR CHAIN_WEIGHTED
BASIS_OF_PROJECTIONS VALUATION PRICES_SECTOR_HARMONIZED_PRICES
LABOR_SECTOR_EMPLOYMENT_TYPE FISCAL_SECTOR_GENERAL_GOVERNMENT_COMPOSITION
FISCAL_SECTOR_VALUATION_OF_DEBT FISCAL_SECTOR_INSTRUMENTS_INCLUDED_IN_GROSS_AND_NET_DEBT
TRADE_SECTOR_OIL_COVERAGE PRIMARY_DOMESTIC_CURRENCY""".split()
METADATA_FIELDS = (
    "COUNTRY",
    "INDICATOR",
    "SCALE",
    "UNIT",
    "LATEST_ACTUAL_ANNUAL_DATA",
    "BASE_YEAR",
    "PRIMARY_DOMESTIC_CURRENCY",
    "METHODOLOGY.ID",
    "METHODOLOGY_NOTES",
    "HISTORICAL_DATA_SOURCE",
    "PUBLICATION_DATE",
    "UPDATE_DATE",
    "START_END_MONTHS_OF_REPORTING_YEAR",
    "FISCAL_SECTOR_GENERAL_GOVERNMENT_COMPOSITION",
    "FISCAL_SECTOR_VALUATION_OF_DEBT",
    "FISCAL_SECTOR_INSTRUMENTS_INCLUDED_IN_GROSS_AND_NET_DEBT",
    "BASIS_OF_PROJECTIONS",
    "VALUATION",
)
# Frozen raw CSV schema; optional SDMX attributes are blank in this representation.
SDMX_HEADERS = [
    "DATAFLOW",
    "COUNTRY",
    "INDICATOR",
    "FREQUENCY",
    "TIME_PERIOD",
    "OBS_VALUE",
    "SCALE",
    "PRECISION",
    "DECIMALS_DISPLAYED",
    "FUNCTIONAL_CAT",
    "INT_ACC_ITEM",
    "NA_STO",
    "GFS_STO",
    "COICOP_1999",
    "TRADE_FLOW",
    "COMMODITY",
    "SOC_CONCEPTS",
    "SECTOR",
    "ACCOUNTING_ENTRY",
    "INDEX_TYPE",
    "PRICES",
    "STATISTICAL_MEASURES",
    "EXRATE",
    "TRANSFORMATION",
    "UNIT",
    "REPORTING_PERIOD_TYPE",
    "DERIVATION_TYPE",
    "OVERLAP",
    "COUNTRY_UPDATE_DATE",
    "DOI",
    "FULL_DESCRIPTION",
    "AUTHOR",
    "PUBLISHER",
    "DEPARTMENT",
    "CONTACT_POINT",
    "TOPIC",
    "TOPIC_DATASET",
    "KEYWORDS",
    "KEYWORDS_DATASET",
    "LANGUAGE",
    "PUBLICATION_DATE",
    "UPDATE_DATE",
    "METHODOLOGY",
    "METHODOLOGY_NOTES",
    "ACCESS_SHARING_LEVEL",
    "ACCESS_SHARING_NOTES",
    "SECURITY_CLASSIFICATION",
    "SOURCE",
    "SHORT_SOURCE_CITATION",
    "FULL_SOURCE_CITATION",
    "LICENSE",
    "SUGGESTED_CITATION",
    "KEY_INDICATOR",
    "SERIES_NAME",
    "LATEST_ACTUAL_ANNUAL_DATA",
    "HISTORICAL_DATA_SOURCE",
    "BASE_YEAR",
    "START_END_MONTHS_OF_REPORTING_YEAR",
    "CHAIN_WEIGHTED",
    "BASIS_OF_PROJECTIONS",
    "VALUATION",
    "PRICES_SECTOR_HARMONIZED_PRICES",
    "LABOR_SECTOR_EMPLOYMENT_TYPE",
    "FISCAL_SECTOR_GENERAL_GOVERNMENT_COMPOSITION",
    "FISCAL_SECTOR_VALUATION_OF_DEBT",
    "FISCAL_SECTOR_INSTRUMENTS_INCLUDED_IN_GROSS_AND_NET_DEBT",
    "TRADE_SECTOR_OIL_COVERAGE",
    "PRIMARY_DOMESTIC_CURRENCY",
]


def require(ok, message):
    if not ok:
        raise ValueError(message)


def canonical(payload):
    """The existing numerical input hash format, without a trailing newline."""
    return json.dumps(
        payload, sort_keys=True, separators=(",", ":"), allow_nan=False
    ).encode()


def write_json(path, value):
    path.write_bytes(canonical(value) + b"\n")


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def tree_hash(path):
    files = sorted(p for p in path.rglob("*") if p.is_file())
    require(not any(p.is_symlink() for p in path.rglob("*")), "Symlink in frozen input")
    body = "".join(f"{sha(p)}  {p.relative_to(path).as_posix()}\n" for p in files)
    return hashlib.sha256(body.encode()).hexdigest()


def input_body(payload):
    require(
        set(payload) == set(INPUT_KEYS) | {"horizonPolicy"},
        "Unexpected inherited payload keys",
    )
    return {key: payload[key] for key in INPUT_KEYS}


def input_hash(body):
    require(set(body) == set(INPUT_KEYS), "Numerical input must have exactly six keys")
    return hashlib.sha256(canonical(body)).hexdigest()


def source_number(token):
    if token == "":
        return None
    require(bool(NUMBER.fullmatch(token)), f"Malformed source numeric token: {token!r}")
    value = float(token)
    require(math.isfinite(value), "Nonfinite source numeric token")
    return value


def workbook_evidence(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        require(
            workbook.sheetnames
            == [
                "October 2026 WEO",
                "Countries",
                "Country Groups",
                "Commodity Prices",
                "Country Group Composition",
            ],
            "Workbook sheets changed",
        )
        title = workbook.worksheets[0]
        require(title["B8"].value == "October 2026", "Wrong workbook vintage")
        date = title["I3"].value
        require(
            isinstance(date, datetime) and date.isoformat() == "2026-10-12T00:00:00",
            "Wrong or missing title date",
        )
        iterator = workbook["Countries"].iter_rows(values_only=True)
        headers = list(next(iterator))
        expected = WORKBOOK_PREFIX + list(range(1980, 2032))
        require(
            len(headers) == len(expected)
            and all(type(a) is type(b) and a == b for a, b in zip(headers, expected)),
            "Workbook header schema changed",
        )
        selected = {}
        for values in iterator:
            row = dict(zip(headers, values))
            if row["INDICATOR.ID"] not in INDICATORS:
                continue
            key = row["COUNTRY.ID"], row["INDICATOR.ID"]
            require(key not in selected, "Duplicate workbook country/indicator")
            require(
                row["DATASET"] == "IMF.RES:WEO(9.0.0)" and row["FREQUENCY"] == "Annual",
                "Workbook dataset/frequency changed",
            )
            for field in ("PUBLICATION_DATE", "UPDATE_DATE"):
                require(
                    row[field] == "2026-10-13T02:00:00Z",
                    "Workbook release metadata changed",
                )
            for year in range(1980, 2032):
                value = row[year]
                require(
                    value is None
                    or value == ""
                    or (type(value) in (int, float) and math.isfinite(value)),
                    "Malformed workbook observation",
                )
                row[year] = None if value == "" else value
            present = any(row[y] is not None for y in range(1980, 2032))
            units = (
                ("Units", "Index")
                if key[1] == "NGDP_D"
                else ("Billions", "Domestic currency")
            )
            require(
                (row["SCALE"], row["UNIT"]) == units
                or (not present and row["SCALE"] is None and row["UNIT"] is None),
                "Workbook unit/scale mismatch",
            )
            actual = row["LATEST_ACTUAL_ANNUAL_DATA"]
            require(
                actual is None or type(actual) in (int, str),
                "Malformed latest-actual metadata",
            )
            selected[key] = row
        codes = sorted({code for code, _ in selected})
        require(
            len(codes) == 197 and len(selected) == len(codes) * len(INDICATORS),
            "Workbook required country/indicator roster incomplete",
        )
        mapped = [config.WEO_CODE_TO_ISO3.get(code, code) for code in codes]
        require(len(mapped) == len(set(mapped)), "Post-alias country collision")
        metadata = [
            {
                "sourceCode": code,
                "iso3c": config.WEO_CODE_TO_ISO3.get(code, code),
                "indicator": indicator,
                **{key: selected[code, indicator][key] for key in METADATA_FIELDS},
            }
            for code in codes
            for indicator in INDICATORS
        ]
        return selected, codes, metadata
    finally:
        workbook.close()


def dense_csv(raw, destination, workbook_rows, codes):
    """Preserve raw numeric strings and units; insert only corroborated blanks."""
    records, seen, placeholders, groups = {}, set(), [], set()
    with raw.open(newline="") as stream:
        reader = csv.DictReader(stream)
        require(reader.fieldnames == SDMX_HEADERS, "SDMX header schema changed")
        for row in reader:
            require(
                None not in row and all(value is not None for value in row.values()),
                "Ragged SDMX row",
            )
            code, indicator = row["COUNTRY"], row["INDICATOR"]
            require(
                indicator in INDICATORS
                and row["FREQUENCY"] == "A"
                and row["DATAFLOW"] == "IMF.RES:WEO(9.0.0)",
                "SDMX series identity changed",
            )
            require(
                code in codes or code in config.WEO_DROP_CODES, "Unknown source country"
            )
            value = source_number(row["OBS_VALUE"])
            if code in config.WEO_DROP_CODES:
                groups.add(code)
                continue
            if row["TIME_PERIOD"] == "":
                require(
                    code == "VEN" and indicator == "NGDP_R" and value is None,
                    "Unexpected empty-time SDMX row",
                )
                placeholders.append((code, indicator))
                continue
            require(
                re.fullmatch(r"\d{4}", row["TIME_PERIOD"]),
                "Malformed annual time token",
            )
            year = int(row["TIME_PERIOD"])
            key = code, indicator, year
            require(key not in seen, "Duplicate raw country/indicator/year")
            seen.add(key)
            if 2000 <= year <= 2031:
                records[key] = row
    require(
        groups == config.WEO_DROP_CODES and placeholders == [("VEN", "NGDP_R")],
        "Excluded aggregate/placeholder roster changed",
    )
    counts, branches, exceptions = Counter(), Counter(), []
    with destination.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=SDMX_HEADERS, lineterminator="\n")
        writer.writeheader()
        for code in codes:
            for indicator in INDICATORS:
                for year in range(2000, 2032):
                    key = code, indicator, year
                    row = records.get(key)
                    expected = workbook_rows[code, indicator][year]
                    value = source_number(row["OBS_VALUE"]) if row else None
                    if value is None:
                        require(
                            expected is None,
                            f"API missing numeric workbook cell: {key}",
                        )
                        counts["missing_api_key_workbook_null"] += 1
                    else:
                        require(
                            expected is not None, f"API numeric/workbook null: {key}"
                        )
                        normalized = value if indicator == "NGDP_D" else value / 1e9
                        delta = abs(normalized - expected)
                        if delta <= 0.00050001:
                            branches["absolute"] += 1
                        else:
                            require(
                                delta <= abs(normalized) * 1e-14,
                                f"Unexplained API/workbook numeric difference: {key}",
                            )
                            branches["relative_only"] += 1
                            exceptions.append(
                                {
                                    "sourceCode": code,
                                    "indicator": indicator,
                                    "year": year,
                                    "api": normalized,
                                    "workbook": expected,
                                    "absoluteDelta": delta,
                                    "unit": "Index"
                                    if indicator == "NGDP_D"
                                    else "Billions of domestic currency",
                                }
                            )
                        counts["numeric_numeric"] += 1
                    if row is None:
                        row = {field: "" for field in SDMX_HEADERS}
                        row.update(
                            DATAFLOW="IMF.RES:WEO(9.0.0)",
                            COUNTRY=code,
                            INDICATOR=indicator,
                            FREQUENCY="A",
                            TIME_PERIOD=str(year),
                        )
                    writer.writerow(row)
    require(
        counts == {"numeric_numeric": 48701, "missing_api_key_workbook_null": 1731}
        and branches == {"absolute": 48693, "relative_only": 8},
        "Source reconciliation counts changed",
    )
    return {
        "denseRows": 50432,
        "counts": dict(counts),
        "precisionBranches": dict(branches),
        "precisionExceptions": exceptions,
        "oneSidedDiscrepancies": [],
        "numericStringsPreserved": 48701,
        "blankObservationsInserted": 1731,
        "excludedAggregates": sorted(groups),
        "excludedEmptyTime": placeholders,
        "denseCsvSha256": sha(destination),
    }


def resolve_release_horizon(body):
    horizon = resolve_horizon(body)
    if horizon["coverageStatus"] == "full" and horizon["weoMaxYear"] < 2031:
        horizon["coverageStatus"] = "shorter"
        horizon["coverageReason"] = (
            f"Country WEO inputs end at {horizon['weoMaxYear']}; "
            "this release extends to 2031 for other countries."
        )
    return horizon


def build(source_dir, output, root=None):
    root = (root or config.repo_root()).resolve()
    output = output.resolve()
    old = {name: root / "data/vintages" / name for name in OLD_TREES}
    require(
        not any(
            output == path.resolve()
            or output.is_relative_to(path.resolve())
            or path.resolve().is_relative_to(output)
            for path in old.values()
        ),
        "Existing vintages and their parent directories are immutable",
    )
    require(
        not output.exists(), "Output already exists; choose a new empty destination"
    )
    for name, path in old.items():
        require(
            tree_hash(path) == OLD_TREES[name], f"Frozen prior vintage changed: {name}"
        )
    for name, expected in SOURCE_FILES.items():
        require(
            (source_dir / name).is_file() and sha(source_dir / name) == expected,
            f"Frozen source identity mismatch: {name}",
        )
    receipts = json.loads((source_dir / "source_receipts.json").read_text())
    require(
        receipts["api"]["sha256"] == SOURCE_FILES["weo_oct2026_raw.csv"]
        and receipts["api"]["url"] == fetch.weo_data_url()
        and receipts["workbook"]["sha256"] == SOURCE_FILES["WEOOct2026all.xlsx"],
        "Source receipt identity mismatch",
    )
    require(
        set(receipts) == {"api", "workbook", "country_codelist"},
        "Unexpected source receipt keys",
    )
    prior_dir = old[RETAINED_FROM]
    prior_manifest = json.loads((prior_dir / "manifest.json").read_text())
    retained_source_receipts = {
        key: {
            field: prior_manifest["raw_inputs"][key][field]
            for field in ("file", "bytes", "sha256")
        }
        for key in ("wpp_medium", "wpp_variants")
    }
    prior_index = json.loads((prior_dir / "json/index.json").read_text())
    entries = prior_index["countries"]
    require(
        len(entries) == 175 and len({r["iso3c"] for r in entries}) == 175,
        "Retained release roster incomplete",
    )
    require(
        all(set(entry) == {"iso3c", "country"} for entry in entries),
        "Unexpected inherited index entry keys",
    )
    workbook_rows, codes, metadata = workbook_evidence(
        source_dir / "WEOOct2026all.xlsx"
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix=".october-2026-", dir=output.parent))
    try:
        dense = stage / "dense.csv"
        ledger = dense_csv(
            source_dir / "weo_oct2026_raw.csv", dense, workbook_rows, codes
        )
        macro = emit.normalise_non_finite(
            weo.build_macrofiscal(
                dense,
                source_dir / "cl_country.json",
                {r["iso3c"]: r["country"] for r in entries},
                year_max=2031,
            )
        )
        require(list(macro.shape) == [6107, 24], "Dense macrofiscal shape changed")
        grouped = {}
        for row in macro.to_dicts():
            grouped.setdefault(row["iso3c"], []).append(row)
        require(
            all(
                [row["years"] for row in rows] == list(range(2001, 2032))
                for rows in grouped.values()
            ),
            "Dense calendar sequence changed",
        )
        dense.unlink()
        json_dir = stage / "json"
        json_dir.mkdir()
        coverage, retained_hashes = [], {}
        for entry in entries:
            iso = entry["iso3c"]
            inherited = json.loads((prior_dir / "json" / f"{iso}.json").read_text())
            body = input_body(inherited)
            require(
                body["iso3c"] == iso and body["country"] == entry["country"],
                "Retained identity mismatch",
            )
            retained_hashes[iso] = {
                key: hashlib.sha256(canonical(body[key])).hexdigest()
                for key in ("demography", "productivity", "climate")
            }
            body["macrofiscal"] = grouped[iso]
            for row in body["macrofiscal"]:
                row["country"] = body["country"]
            horizon = resolve_release_horizon(body)
            policy = {
                "id": POLICY,
                "dataRevision": REVISION,
                "sourceVintage": SOURCE_VINTAGE,
                **horizon,
                "inputSha256": input_hash(body),
            }
            write_json(json_dir / f"{iso}.json", {**body, "horizonPolicy": policy})
            previous = inherited["horizonPolicy"]
            coverage.append(
                {
                    "iso3c": iso,
                    "country": body["country"],
                    **policy,
                    "previousCoverageStatus": previous["coverageStatus"],
                    "previousWeoMaxYear": previous["weoMaxYear"],
                    "coverageChanged": previous["coverageStatus"]
                    != horizon["coverageStatus"]
                    or previous["weoMaxYear"] != horizon["weoMaxYear"],
                }
            )
        verified_hashes = {
            p.stem: hashlib.sha256(canonical(json.loads(p.read_text()))).hexdigest()
            for p in sorted((old["weo-2024-10"] / "json").glob("*.json"))
            if p.stem != "index"
        }
        require(
            verified_hashes == prior_index["verifiedInputHashes"],
            "Verified input identities changed",
        )
        index = {
            "vintage": SOURCE_VINTAGE,
            "label": LABEL,
            "count": len(entries),
            "countries": entries,
            "dataRevision": REVISION,
            "calculationPolicy": POLICY,
            "releaseWeoMaxYear": 2031,
            "verifiedInputHashes": verified_hashes,
        }
        write_json(json_dir / "index.json", index)
        source_metadata = {
            "sourceWorkbookSha256": SOURCE_FILES["WEOOct2026all.xlsx"],
            "sourceVintage": SOURCE_VINTAGE,
            "schema": list(METADATA_FIELDS),
            "records": metadata,
            "note": (
                "Raw per-series metadata; H is not an actual-outturn cutoff. "
                "Fiscal-year strings and nulls are preserved."
            ),
        }
        write_json(stage / "source_metadata.json", source_metadata)
        ledger.update(
            {
                "sourceVintage": SOURCE_VINTAGE,
                "metadataSha256": sha(stage / "source_metadata.json"),
                "retainedSliceSha256": retained_hashes,
            }
        )
        write_json(stage / "source_validation.json", ledger)
        inventory = {
            p.relative_to(stage).as_posix(): sha(p)
            for p in sorted(stage.rglob("*"))
            if p.is_file()
        }
        require(
            set(p.name for p in json_dir.iterdir())
            == {f"{r['iso3c']}.json" for r in entries} | {"index.json"},
            "Release JSON inventory is not exactly the complete roster plus index",
        )
        manifest = {
            "vintage": SOURCE_VINTAGE,
            "sourceVintage": SOURCE_VINTAGE,
            "label": LABEL,
            "dataRevision": REVISION,
            "calculationPolicy": POLICY,
            "retained_from": RETAINED_FROM,
            "base_vintage": "weo-2024-10",
            "raw_inputs": {
                name: {
                    "file": name,
                    "bytes": (source_dir / name).stat().st_size,
                    "sha256": expected,
                }
                for name, expected in SOURCE_FILES.items()
            },
            "source_receipts": receipts,
            "retained_source_receipts": retained_source_receipts,
            "macrofiscal_year_max": 2031,
            "macrofiscal_rows": 6107,
            "retained_macrofiscal_rows": 5425,
            "country_json": {
                "files": 175,
                "bytes": sum(
                    (json_dir / f"{r['iso3c']}.json").stat().st_size for r in entries
                ),
            },
            "selectable_countries": 175,
            "coverage": coverage,
            "coverageTransitions": [row for row in coverage if row["coverageChanged"]],
            "coverage_counts": dict(Counter(r["coverageStatus"] for r in coverage)),
            "retained_input_slices_sha256": retained_hashes,
            "previous_vintage_tree_sha256": OLD_TREES,
            "source_metadata": {
                "file": "source_metadata.json",
                "sha256": sha(stage / "source_metadata.json"),
            },
            "source_validation": {
                "file": "source_validation.json",
                "sha256": sha(stage / "source_validation.json"),
            },
            "generated_files_sha256": inventory,
            "inputHashDefinition": (
                "SHA256 of compact sorted-key ASCII JSON, allow_nan=False, no "
                "newline, exact six-key body "
                "iso3c/country/macrofiscal/demography/productivity/climate; "
                "horizonPolicy excluded."
            ),
            "observationStatus": (
                "API dataonly attributes are blank; workbook raw latest-actual "
                "metadata is provided separately per country/indicator. H is "
                "the usable modeling boundary."
            ),
            "releaseIdentity": (
                "Rolling dataflow version 9.0.0 is not a vintage pin. Frozen "
                "API values/nulls exhaustively corroborate the independently "
                "titled/published October full workbook to its presentation "
                "precision."
            ),
            "sources": {
                "macrofiscal": {
                    "provider": "IMF World Economic Outlook",
                    "vintage": "October 2026",
                    "endpoint": receipts["api"]["url"],
                    "indicators": dict(config.WEO_INDICATORS),
                },
                "demography": {
                    "provider": "UN World Population Prospects",
                    "vintage": "2024 revision",
                    "retainedFrom": RETAINED_FROM,
                },
                "productivity": {
                    "provider": "carried forward",
                    "vintage": "weo-2024-10",
                },
                "climate": {"provider": "carried forward", "vintage": "weo-2024-10"},
            },
            "generator": {
                "module": "qcraft_pipeline.october_2026",
                "sha256": sha(Path(__file__)),
            },
            "rebuild": "python -m qcraft_pipeline.october_2026 (offline)",
        }
        write_json(stage / "manifest.json", manifest)
        require(
            manifest["coverage_counts"]
            == {"full": 161, "shorter": 6, "unsupported": 8},
            "Unexpected October coverage outcome",
        )
        for name, path in old.items():
            require(
                tree_hash(path) == OLD_TREES[name],
                "Existing vintage changed during build",
            )
        require(not output.exists(), "Output appeared during build; refusing overwrite")
        stage.rename(output)
        return manifest
    except BaseException:
        shutil.rmtree(stage)
        raise


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-dir", type=Path, default=config.cache_dir() / "october-2026"
    )
    parser.add_argument("--out", type=Path, default=config.vintage_dir(REVISION))
    args = parser.parse_args(argv)
    manifest = build(args.source_dir, args.out)
    print(
        json.dumps(
            {"revision": REVISION, "coverage": manifest["coverage_counts"]}, indent=2
        )
    )


if __name__ == "__main__":
    main()
